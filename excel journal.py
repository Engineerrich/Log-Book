import datetime
import os
from openpyxl import Workbook, load_workbook

FILENAME = r"C:\Users\rrobb\Documents\my_journal.xls"

def save_to_excel():
    # 1. Get current data
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    mood = input("How are you feeling (1-10)? ")
    thought = input("What's on your mind? ")

    # 2. Open existing file or create a new one
    if os.path.exists(FILENAME):
        wb = load_workbook(FILENAME)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Mood Log"
        # Create headers for the first time
        ws.append(["Timestamp", "Mood Rating", "Thought"])

    # 3. Append the new row of data
    ws.append([now, mood, thought])

    # 4. Save the file
    wb.save(FILENAME)
    print(f"\n✅ Entry saved to {os.path.abspath(FILENAME)}")

if __name__ == "__main__":
    save_to_excel()

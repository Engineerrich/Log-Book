import datetime
import os
from openpyxl import Workbook, load_workbook

FILENAME = r"C:\Users\Jennifer Robb\Documents\RICH's Files\Python Programs\Mileage_Journal.xlsx"

def save_to_excel():
    # 1. Get current data
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    truck = input("Truck C,P,D: ")
    location = input("Where are you? ")
    turnonmileage = input("What is your turn on mileage? ")
    turnoffmileage = input("What is your turn off mileage? ")

    # 2. Open existing file or create a new one
    if os.path.exists(FILENAME):
        wb = load_workbook(FILENAME)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Mileage Log"
        # Create headers for the first time
        ws.append(["Timestamp", "Truck", "Location", "Turn On Mileage", "Turn Off Mileage"])

    # 3. Append the new row of data
    ws.append([now, truck, location, turnonmileage, turnoffmileage])

    # 4. Save the file
    wb.save(FILENAME)
    print(f"\n✅ Entry saved to {os.path.abspath(FILENAME)}")

if __name__ == "__main__":
    save_to_excel()
